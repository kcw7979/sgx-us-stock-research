"""Deterministic synthetic TEST data only; never imported by application code."""
import unittest
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from equity_research.export import excel_export
from equity_research.indicators import indicators
from equity_research.portfolio import allocate, position_size


class IndicatorTests(unittest.TestCase):
    def bars(self, prices):
        return pd.DataFrame({"adj_close": np.asarray(prices, dtype=float)})

    def test_linear_series_windows_and_no_lookahead(self):
        result = indicators(self.bars(range(1, 301)))
        self.assertTrue(result.ma50.iloc[:49].isna().all())
        self.assertEqual(result.ma50.iloc[49], 25.5)
        self.assertTrue(result.ma200.iloc[:199].isna().all())
        self.assertEqual(result.ma200.iloc[199], 100.5)
        self.assertTrue(result.high_52w.iloc[:251].isna().all())
        self.assertEqual(result.high_52w.iloc[-1], 300)
        self.assertEqual(result.low_52w.iloc[-1], 49)
        self.assertEqual(result.rsi14.iloc[14], 100)
        self.assertTrue(result.rsi14.iloc[:14].isna().all())
        prefix = indicators(self.bars(range(1, 251)))
        pd.testing.assert_frame_equal(result.iloc[:250], prefix)

    def test_returns_and_annualized_sample_volatility(self):
        prices = np.array([100., 101., 99., 103., 100.] * 10)
        result = indicators(self.bars(prices))
        expected = prices[1:] / prices[:-1] - 1
        np.testing.assert_allclose(result.daily_return.iloc[1:], expected)
        self.assertTrue(np.isnan(result.daily_return.iloc[0]))
        self.assertTrue(result.volatility_20d.iloc[:20].isna().all())
        self.assertAlmostEqual(result.volatility_20d.iloc[20], np.std(expected[:20], ddof=1) * np.sqrt(252))

    def test_rsi_wilder_reference(self):
        # Published-style hand-checkable arithmetic seed, then one recursive step.
        prices = [100, 102, 101, 104, 102, 106, 103, 108, 104, 110, 105, 112, 106, 114, 107, 116]
        result = indicators(self.bars(prices))
        gain, loss = 35 / 14, 28 / 14
        self.assertAlmostEqual(result.rsi14.iloc[14], 100 - 100 / (1 + gain / loss))
        gain, loss = (gain * 13 + 9) / 14, loss * 13 / 14
        self.assertAlmostEqual(result.rsi14.iloc[15], 100 - 100 / (1 + gain / loss))

    def test_flat_and_falling_rsi(self):
        self.assertEqual(indicators(self.bars([10] * 30)).rsi14.iloc[-1], 50)
        self.assertEqual(indicators(self.bars(range(30, 0, -1))).rsi14.iloc[-1], 0)

    def test_empty_and_short_history_keep_unavailable_indicators(self):
        self.assertTrue(indicators(self.bars([])).empty)
        frame = indicators(self.bars([10, 11]))
        self.assertTrue(frame[["ma50", "ma200", "rsi14", "high_52w", "low_52w", "volatility_20d"]].isna().all().all())


class PortfolioTests(unittest.TestCase):
    def test_allocation_preserves_cash(self):
        frame = allocate(10000, {"S63": .3, "NVDA": .5})
        self.assertAlmostEqual(frame.base_amount.sum(), 10000)
        self.assertAlmostEqual(frame.iloc[-1].base_amount, 2000)

    def test_allocation_rejects_bad_inputs(self):
        for capital, weights in [(0, {"A": 1}), (100, {}), (100, {"A": 1.1}),
                                 (100, {"A": -.1}), (100, {"A": np.nan}),
                                 (np.inf, {"A": 1})]:
            with self.subTest(capital=capital, weights=weights), self.assertRaises(ValueError):
                allocate(capital, weights)

    def test_risk_fx_and_board_lot_sizing(self):
        result = position_size(100000, .01, 50, 45, 1.35, .3, 100)
        self.assertEqual(result["shares"], 100)
        self.assertAlmostEqual(result["risk_base"], 675)
        self.assertLessEqual(result["risk_base"], result["risk_budget_base"])
        self.assertLessEqual(result["cost_base"], 30000)

    def test_cash_cap_and_zero_share_result(self):
        self.assertEqual(position_size(1000, .1, 100, 99, 1, .2)["shares"], 2)
        self.assertEqual(position_size(100, .01, 50, 45, 1, 1, 100)["shares"], 0)

    def test_reject_invalid_sizing_inputs(self):
        defaults = dict(capital=1000, risk_fraction=.01, entry=10, stop=9,
                        fx_to_base=1, allocation_fraction=.5, lot_size=1)
        for change in [{"entry": 0}, {"stop": 10}, {"stop": 11}, {"stop": -1},
                       {"fx_to_base": 0}, {"capital": np.inf}, {"risk_fraction": np.nan},
                       {"risk_fraction": 0}, {"allocation_fraction": 1.1}, {"lot_size": 1.5},
                       {"lot_size": True}, {"lot_size": 0}]:
            with self.subTest(change=change), self.assertRaises(ValueError):
                position_size(**(defaults | change))


class ExportTests(unittest.TestCase):
    def test_excel_roundtrip_includes_provenance_and_indicators(self):
        frame = pd.DataFrame({
            "date": pd.date_range("2025-01-01", periods=260),
            "adj_close": np.arange(1., 261.), "currency": "SGD",
            "provider": "MOCK TEST DATA ONLY",
            "retrieved_at": pd.Timestamp("2026-01-01", tz="UTC"),
            "source_timestamp": pd.date_range("2025-01-01", periods=260, tz="UTC"),
        })
        workbook = load_workbook(BytesIO(excel_export({"S63": frame}, allocate(100, {"S63": .5}))))
        self.assertEqual(workbook.sheetnames, ["Read me", "S63", "Allocation"])
        headers = [cell.value for cell in workbook["S63"][1]]
        self.assertIn("ma200", headers)
        self.assertIn("currency", headers)
        self.assertIn("provider", headers)
        timestamp = workbook["S63"].cell(2, headers.index("retrieved_at") + 1).value
        self.assertTrue(timestamp.endswith("+00:00"))
        self.assertEqual(workbook["S63"].max_row, 261)
        self.assertEqual(workbook["S63"].freeze_panes, "A2")


if __name__ == "__main__":
    unittest.main()
